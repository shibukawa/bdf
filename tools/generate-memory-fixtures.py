#!/usr/bin/env python3
"""Generate larger XLSX/PPTX inputs for the converter memory benchmark.

Requires xlsxwriter and python-pptx. The fixtures exercise the multi-sheet
and multi-slide release paths introduced in fa64bb7. ZIP metadata can vary
between runs; the benchmark JSON records the hashes of the measured files.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import xlsxwriter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


def workbook(path: Path) -> None:
    book = xlsxwriter.Workbook(str(path))
    book.set_properties({"title": "Synthetic memory benchmark", "author": "bdf benchmark"})
    heading = book.add_format({"bold": True, "bg_color": "#DDEBF7"})
    for sheet_index in range(4):
        sheet = book.add_worksheet(f"Sheet {sheet_index + 1}")
        sheet.freeze_panes(1, 1)
        sheet.set_column(0, 7, 14)
        for column in range(8):
            sheet.write(0, column, f"Column {column + 1}", heading)
        for row in range(1, 3001):
            sheet.write(row, 0, f"Item {sheet_index}-{row:04d}")
            for column in range(1, 8):
                sheet.write_number(row, column, (row * (column + 3) + sheet_index) % 100000)
    book.close()


def presentation(path: Path) -> None:
    deck = Presentation()
    deck.slide_width = Inches(10)
    deck.slide_height = Inches(5.625)
    for slide_index in range(100):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for shape_index in range(36):
            x = Inches(0.2 + (shape_index % 9) * 1.08)
            y = Inches(0.25 + (shape_index // 9) * 1.28)
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(1.0), Inches(0.9)
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(
                (slide_index * 3 + shape_index * 7) % 255,
                (slide_index * 5 + shape_index * 11) % 255,
                (slide_index * 13 + shape_index * 17) % 255,
            )
            shape.text = (
                f"Slide {slide_index + 1} block {shape_index + 1}: "
                "reproducible benchmark content"
            )
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)
    deck.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True, help="new output directory")
    output = parser.parse_args().output.expanduser().resolve()
    output.mkdir(parents=True, exist_ok=False)
    workbook(output / "many-sheets.xlsx")
    presentation(output / "many-slides.pptx")
    for path in output.iterdir():
        print(f"{path}: {path.stat().st_size} bytes")


if __name__ == "__main__":
    main()
