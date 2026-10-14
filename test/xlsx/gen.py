"""Generates the Excel test workbooks in converter/xlsx/testdata with openpyxl.

Usage: python3 test/xlsx/gen.py   (pip install openpyxl pillow)
basic.xlsx exercises cell formats: number formats, fonts and rich text,
fills, borders, alignment (wrapping Latin and Japanese text, overflow,
rotation, indent, shrink to fit), merged cells, row and column sizes,
frozen panes, notes and hyperlinks, a hidden sheet, and a sheet large
enough to span several tiles. features.xlsx exercises conditional formats,
tables, a chart, a picture, a text box shape and a chart sheet. math.xlsx
holds equation shapes as Excel writes them: Office Math (a14:m) in the
mc:Choice of an mc:AlternateContent whose fallback holds the formula as
text, on a line of its own and inline in text. print.xlsx holds sheets with
page setups, for the print view. Names of the functions as arguments write
only their workbooks (python3 test/xlsx/gen.py printing).
"""
import datetime
import io
import os
import re
import struct
import sys
import zipfile
import zlib

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chartsheet import Chartsheet
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import (CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, IconSetRule, Rule)
from openpyxl.styles import Alignment, Border, Font, GradientFill, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(ROOT, "converter", "xlsx", "testdata")
FIXED = datetime.datetime(2026, 9, 1, 9, 0, 0)


def png(w, h, fn):
    """A small RGB PNG from fn(x, y) -> (r, g, b)."""
    raw = b""
    for y in range(h):
        raw += b"\x00" + b"".join(bytes(fn(x, y)) for x in range(w))

    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw, 9)) + chunk(b"IEND", b""))


def props(wb, title):
    wb.properties.title = title
    wb.properties.creator = "bdf test"
    wb.properties.lastModifiedBy = "bdf test"
    wb.properties.created = FIXED
    wb.properties.modified = FIXED


thin = Side(style="thin", color="000000")
medium = Side(style="medium", color="1F4E79")


def basic():
    wb = Workbook()
    props(wb, "Excel test workbook")
    ws = wb.active
    ws.title = "Sales"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 13
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 8
    ws.column_dimensions["H"].hidden = True

    # title: merged, bold, filled
    ws.merge_cells("A1:F1")
    ws["A1"] = "Quarterly sales 数字の表"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    head = ["Region", "Q1", "Q2", "Growth", "Updated", "Flag"]
    for i, h in enumerate(head):
        c = ws.cell(row=3, column=i + 1, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
        c.border = Border(top=thin, bottom=medium, left=thin, right=thin)
        c.alignment = Alignment(horizontal="center")
    rows = [
        ("North ノース", 125000, 131250.5, 0.05, datetime.date(2026, 3, 31), True),
        ("South サウス", 98000, 91140, -0.07, datetime.date(2026, 6, 30), False),
        ("East", 1234567.891, 1300000, 0.053, datetime.date(2026, 9, 30), True),
        ("West", 0, 0.25, 1.5, datetime.datetime(2026, 12, 31, 18, 30), False),
    ]
    for r, row in enumerate(rows, start=4):
        for i, v in enumerate(row):
            c = ws.cell(row=r, column=i + 1, value=v)
            c.border = Border(left=thin, right=thin, bottom=Side(style="hair", color="808080"))
        ws.cell(row=r, column=2).number_format = "#,##0"
        ws.cell(row=r, column=3).number_format = '#,##0.00;[Red]-#,##0.00'
        ws.cell(row=r, column=4).number_format = "0.0%"
        ws.cell(row=r, column=5).number_format = "yyyy-mm-dd"
    ws["E7"].number_format = "yyyy/m/d h:mm"
    ws["C5"] = -1234.5
    ws["A8"] = "Total"
    ws["A8"].font = Font(bold=True)
    ws["B8"] = "=SUM(B4:B7)"
    ws["B8"].number_format = "#,##0"
    ws["B8"].border = Border(top=thin, bottom=Side(style="double"))
    ws.freeze_panes = "A4"

    # text: overflow, clipping, wrapping (Latin and Japanese), alignment
    ws["A10"] = "This text is long and flows over the empty cells to its right."
    ws["A11"] = "Clipped by the next cell"
    ws["B11"] = "next"
    ws["F12"] = "Right-aligned text flows left"
    ws["F12"].alignment = Alignment(horizontal="right")
    ws["A13"] = "日本語の文章は、単語の区切りに空白を使わないため、文字と文字の間で改行します。句読点「、」や「。」は行頭に来ません。"
    ws["A13"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A13:C13")
    ws.row_dimensions[13].height = 62
    ws["D13"] = "Wrapped Latin text breaks after spaces between words."
    ws["D13"].alignment = Alignment(wrap_text=True, vertical="center")
    ws["E13"] = "Justified text spreads the words of every line but the last to both edges."
    ws["E13"].alignment = Alignment(horizontal="justify", vertical="top")
    ws["A15"] = "indent 2"
    ws["A15"].alignment = Alignment(indent=2)
    ws["B15"] = "shrink to fit this text"
    ws["B15"].alignment = Alignment(shrink_to_fit=True)
    ws["C15"] = "centered across"
    ws["C15"].alignment = Alignment(horizontal="centerContinuous")
    ws["D15"].alignment = Alignment(horizontal="centerContinuous")
    ws["F15"] = "-~"
    ws["F15"].alignment = Alignment(horizontal="fill")
    ws["A16"] = "rotated 45°"
    ws["A16"].alignment = Alignment(text_rotation=45)
    ws["B16"] = "up 90°"
    ws["B16"].alignment = Alignment(text_rotation=90, horizontal="center")
    ws["C16"] = "down"
    ws["C16"].alignment = Alignment(text_rotation=180)
    ws["D16"] = "縦書き"
    ws["D16"].alignment = Alignment(text_rotation=255, horizontal="center")
    ws.row_dimensions[16].height = 60
    # numbers that do not fit, General in a narrow column
    ws["F17"] = 123456789012
    ws["F17"].number_format = "#,##0"
    ws["F18"] = 3.14159265358979
    ws["F19"] = "=1/3"
    # rich text, fonts
    ws["A18"] = CellRichText("Rich ", TextBlock(InlineFont(b=True, color="C00000"), "bold red"), " and ",
                             TextBlock(InlineFont(i=True, u="single", rFont="Times New Roman"), "italic"))
    ws["A19"] = "Strike"
    ws["A19"].font = Font(strike=True, color="7F7F7F")
    ws["B19"] = "H"
    ws["C19"] = "x2"
    ws["C19"].font = Font(vertAlign="superscript")
    ws["A20"] = "Big text"
    ws["A20"].font = Font(size=20, name="Georgia")
    # fills and borders
    ws["A22"] = "gray125"
    ws["A22"].fill = PatternFill("gray125", fgColor="4472C4", bgColor="FFFFFF")
    ws["B22"] = "darkGrid"
    ws["B22"].fill = PatternFill("darkGrid", fgColor="A9D08E", bgColor="FFFFFF")
    ws["C22"] = "gradient"
    ws["C22"].fill = GradientFill(stop=("FFFFFF", "F4B084"), degree=90)
    ws["D22"] = "theme"
    ws["D22"].fill = PatternFill("solid", fgColor="FFE699")
    ws["E22"] = "diagonal"
    ws["E22"].border = Border(diagonal=Side(style="thin", color="FF0000"), diagonalDown=True, diagonalUp=True)
    for col, style in zip("ABCDEF", ["thin", "medium", "thick", "double", "dashed", "dotted"]):
        c = ws[f"{col}24"]
        c.value = style
        c.border = Border(top=Side(style=style), bottom=Side(style=style), left=Side(style=style), right=Side(style=style))
    ws.row_dimensions[21].hidden = True
    ws["A21"] = "hidden row"
    # a note and a hyperlink
    ws["B4"].comment = Comment("Highest quarter", "bdf")
    ws["A26"] = "example.com"
    ws["A26"].hyperlink = "https://example.com/"
    ws["A26"].font = Font(color="0563C1", underline="single")
    ws["A27"] = "Date"
    ws["B27"] = datetime.date(2026, 5, 1)
    ws["B27"].number_format = 'd-mmm-yy'
    ws["C27"] = datetime.date(2026, 5, 1)
    ws["C27"].number_format = "dddd"
    ws["A28"] = 0.375
    ws["A28"].number_format = "# ?/8"
    ws["B28"] = 12345.678
    ws["B28"].number_format = "0.00E+00"
    ws["C28"] = 1234.5
    ws["C28"].number_format = '_("$"* #,##0.00_);_("$"* \\(#,##0.00\\);_("$"* "-"??_);_(@_)'
    ws["D28"] = -1234.5
    ws["D28"].number_format = '_("$"* #,##0.00_);_("$"* \\(#,##0.00\\);_("$"* "-"??_);_(@_)'
    ws["E28"] = "#N/A"
    ws["E28"].data_type = "e"

    hidden = wb.create_sheet("Hidden")
    hidden["A1"] = "Hidden sheet"
    hidden.sheet_state = "hidden"

    big = wb.create_sheet("Big")
    big["A1"] = "Row"
    big["B1"] = "Value"
    for r in range(2, 302):
        big.cell(row=r, column=1, value=f"row {r}")
        big.cell(row=r, column=2, value=r * r)
    # text across a vertical tile boundary (2048 pt = row 137 at 15 pt)
    big["C137"] = "straddles the tile boundary"
    big["C137"].font = Font(size=28)
    big.row_dimensions[137].height = 40
    # text flowing over a horizontal tile boundary (at 2048 pt: column AI with
    # the 60 pt columns of the test fonts)
    big["AG5"] = "flows across the tile boundary to the right"
    wb.save(os.path.join(OUT, "basic.xlsx"))
    fix_core(os.path.join(OUT, "basic.xlsx"))


SHAPE = """<xdr:twoCellAnchor editAs="oneCell"><xdr:from><xdr:col>8</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>31</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>11</xdr:col><xdr:colOff>304800</xdr:colOff><xdr:row>36</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to><xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="100" name="Note box" descr="A rounded note box"/><xdr:cNvSpPr/></xdr:nvSpPr><xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2438400" cy="952500"/></a:xfrm><a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom><a:solidFill><a:schemeClr val="accent1"><a:lumMod val="20000"/><a:lumOff val="80000"/></a:schemeClr></a:solidFill><a:ln w="12700"><a:solidFill><a:schemeClr val="accent1"/></a:solidFill></a:ln></xdr:spPr><xdr:txBody><a:bodyPr vertOverflow="clip" horzOverflow="clip" rtlCol="0" anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="ja-JP" altLang="en-US" sz="1100"/><a:t>図形のテキスト shape text</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>"""


def features():
    wb = Workbook()
    props(wb, "Excel features")
    ws = wb.active
    ws.title = "Features"
    ws["A1"] = "Conditional formats"
    ws["A1"].font = Font(bold=True, size=13)
    vals = [3, 18, 7, 42, 25, 18, 9, 31, 12, 36]
    ws["A2"], ws["B2"], ws["C2"], ws["D2"], ws["E2"] = "scale", "bar", "icons", "cellIs", "top/dup"
    for i, v in enumerate(vals, start=3):
        for col in "ABCDE":
            ws[f"{col}{i}"] = v
    ws.conditional_formatting.add("A3:A12", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50,
                                                          mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    ws.conditional_formatting.add("B3:B12", DataBarRule(start_type="min", end_type="max", color="638EC6"))
    ws.conditional_formatting.add("C3:C12", IconSetRule("3Arrows", "percent", [0, 33, 67]))
    red = DifferentialStyle(font=Font(color="9C0006"), fill=PatternFill(bgColor="FFC7CE"))
    ws.conditional_formatting.add("D3:D12", Rule(type="cellIs", operator="greaterThan", formula=["20"], dxf=red))
    green = DifferentialStyle(font=Font(bold=True, color="006100"), fill=PatternFill(bgColor="C6EFCE"))
    ws.conditional_formatting.add("E3:E12", Rule(type="top10", rank=3, dxf=green))
    yellow = DifferentialStyle(fill=PatternFill(bgColor="FFEB9C"))
    ws.conditional_formatting.add("E3:E12", Rule(type="duplicateValues", dxf=yellow))
    ws.column_dimensions["G"].width = 12
    ws["G2"] = "text rule"
    for i, s in enumerate(["apple", "banana", "cherry", "grape"], start=3):
        ws[f"G{i}"] = s
    ws.conditional_formatting.add("G3:G6", Rule(type="containsText", operator="containsText", text="an",
                                                formula=['NOT(ISERROR(SEARCH("an",G3)))'], dxf=red))

    # a table with totals
    ws["A15"] = "Table"
    ws["A15"].font = Font(bold=True, size=13)
    data = [("Item", "Units", "Price"), ("Pens", 120, 1.5), ("Paper", 40, 4.25), ("Ink", 15, 22.0), ("Staples", 300, 0.02)]
    for r, row in enumerate(data, start=16):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    ws["A21"] = "Total"
    ws["B21"] = "=SUBTOTAL(109,B17:B20)"
    tab = Table(displayName="Supplies", ref="A16:C21", totalsRowCount=1)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)

    # a chart and a picture
    chart = BarChart()
    chart.title = "Units"
    chart.y_axis.title = "units"
    chart.add_data(Reference(ws, min_col=2, min_row=16, max_row=20), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=17, max_row=20))
    chart.width, chart.height = 12, 7
    ws.add_chart(chart, "E15")
    img = Image(io.BytesIO(png(64, 48, lambda x, y: ((x * 4) % 256, (y * 5) % 256, 160 if (x // 8 + y // 8) % 2 else 60))))
    img.anchor = "I3"
    ws.add_image(img)

    cs = wb.create_chartsheet("Chart")
    line = LineChart()
    line.title = "Values"
    line.add_data(Reference(ws, min_col=1, min_row=2, max_row=12), titles_from_data=True)
    cs.add_chart(line)
    wb.save(os.path.join(OUT, "features.xlsx"))
    add_shape(os.path.join(OUT, "features.xlsx"))
    fix_core(os.path.join(OUT, "features.xlsx"))


def fix_core(path):
    """Sets the modified date openpyxl stamps with the time of saving."""
    rewrite(path, {"docProps/core.xml": lambda s: re.sub(r"(<dcterms:modified[^>]*>)[^<]*", r"\g<1>" + FIXED.isoformat() + "Z", s)})


def rewrite(path, edits):
    """Rewrites parts of a package with functions of their text."""
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename in edits:
                data = edits[info.filename](data.decode("utf-8")).encode("utf-8")
            info.date_time = FIXED.timetuple()[:6]
            z.writestr(info, data)
    with open(path, "wb") as f:
        f.write(out.getvalue())


def add_shape(path):
    """Adds a text box shape to the first drawing (openpyxl writes none)."""
    def edit(s):
        s = s.rstrip()
        # openpyxl writes the spreadsheet drawing namespace as the default one
        shape = SHAPE.replace("<xdr:", "<").replace("</xdr:", "</")
        s = re.sub(r"</wsDr>$", lambda m: shape + m.group(0), s)
        return s.replace("<wsDr ", '<wsDr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" ', 1)
    rewrite(path, {"xl/drawings/drawing1.xml": edit})


NS_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def mr(text, sty=None):
    """A math run as Excel writes it: the letters styled (𝑥 for an italic x)."""
    t = text
    if sty != "p":
        t = "".join(chr(0x1D44E + ord(c) - 97) if "a" <= c <= "z" and c != "h" else ("\u210e" if c == "h" else c) for c in text)
    m = f'<m:rPr><m:sty m:val="{sty}"/></m:rPr>' if sty else ""
    return f'<m:r>{m}<a:rPr lang="en-US" sz="1400" b="0" i="1"><a:latin typeface="Cambria Math"/></a:rPr><m:t>{t}</m:t></m:r>'


def equation_anchor(sid, col0, row0, col1, row1, paragraphs, fallback):
    def anchor(body):
        return (f'<xdr:twoCellAnchor><xdr:from><xdr:col>{col0}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{row0}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
                f'<xdr:to><xdr:col>{col1}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{row1}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
                f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{sid}" name="TextBox {sid}"/><xdr:cNvSpPr txBox="1"/></xdr:nvSpPr>'
                f'<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="3000000" cy="600000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
                f'<a:solidFill><a:schemeClr val="lt1"/></a:solidFill><a:ln w="9525"><a:solidFill><a:schemeClr val="lt1"><a:shade val="50000"/></a:schemeClr></a:solidFill></a:ln></xdr:spPr>'
                f'<xdr:txBody><a:bodyPr vertOverflow="clip" horzOverflow="clip" wrap="square" rtlCol="0" anchor="ctr"/><a:lstStyle/>{body}</xdr:txBody></xdr:sp>'
                f'<xdr:clientData/></xdr:twoCellAnchor>')
    return (f'<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
            f'<mc:Choice xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" Requires="a14">{anchor("".join(paragraphs))}</mc:Choice>'
            f'<mc:Fallback>{anchor(fallback)}</mc:Fallback></mc:AlternateContent>')


def math():
    wb = Workbook()
    props(wb, "Equations")
    ws = wb.active
    ws.title = "Equations"
    ws["A1"] = "Equation shapes"
    ws["A1"].font = Font(bold=True, size=13)
    for i, (a, b, c) in enumerate([(1, -3, 2), (1, 2, 1), (2, 1, -1)], start=3):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"], ws[f"D{i}"] = a, b, c, b * b - 4 * a * c
    ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "a", "b", "c", "D"
    img = Image(io.BytesIO(png(16, 16, lambda x, y: (200, 200, 200))))
    img.anchor = "A12"
    ws.add_image(img)
    path = os.path.join(OUT, "math.xlsx")
    wb.save(path)
    quad = (mr("x=") + f'<m:f><m:fPr/><m:num>{mr("−b±")}<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/><m:e>'
            f'<m:sSup><m:e>{mr("b")}</m:e><m:sup>{mr("2")}</m:sup></m:sSup>{mr("−4ac")}</m:e></m:rad></m:num><m:den>{mr("2a")}</m:den></m:f>')
    disc = f'<m:sSup><m:e>{mr("b")}</m:e><m:sup>{mr("2")}</m:sup></m:sSup>{mr("−4ac")}'
    shapes = (equation_anchor(20, 5, 1, 10, 6,
                              [f'<a:p><a14:m><m:oMathPara xmlns:m="{NS_M}"><m:oMathParaPr><m:jc m:val="centerGroup"/></m:oMathParaPr>'
                               f'<m:oMath>{quad}</m:oMath></m:oMathPara></a14:m><a:endParaRPr lang="en-US" sz="1400"/></a:p>'],
                              '<a:p><a:r><a:rPr lang="en-US" sz="1100"/><a:t>FALLBACK TEXT x=(−b±√(b^2−4ac))/2a</a:t></a:r></a:p>')
              + equation_anchor(21, 5, 7, 10, 10,
                                ['<a:p><a:r><a:rPr lang="ja-JP" altLang="en-US" sz="1100"/><a:t>判別式 </a:t></a:r>'
                                 f'<a14:m><m:oMath xmlns:m="{NS_M}">{mr("D=")}{disc}</m:oMath></a14:m>'
                                 '<a:r><a:rPr lang="ja-JP" altLang="en-US" sz="1100"/><a:t> が正なら実数解は二つ。</a:t></a:r></a:p>'],
                                '<a:p><a:r><a:rPr lang="en-US" sz="1100"/><a:t>FALLBACK TEXT D=b^2−4ac</a:t></a:r></a:p>'))

    def edit(xml):
        xml = xml.rstrip()
        shape = shapes.replace("<xdr:", "<").replace("</xdr:", "</")
        xml = re.sub(r"</wsDr>$", lambda m: shape + m.group(0), xml)
        return xml.replace("<wsDr ", '<wsDr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" ', 1)
    rewrite(path, {"xl/drawings/drawing1.xml": edit})
    fix_core(path)


def printing():
    """print.xlsx: sheets with page setups, for the print view (-param
    views=print): print areas and titles, paper, margins, scaling and fitting,
    manual page breaks, headers and footers, gridlines and headings."""
    wb = Workbook()
    props(wb, "Printed sheets")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill("solid", fgColor="DDEBF7")

    # a long table on A4: titles on every page, a manual break, gridlines
    ws = wb.active
    ws.title = "Report"
    for col, w in zip("ABCDEF", (6, 18, 7, 10, 12, 13)):
        ws.column_dimensions[col].width = w
    ws["A1"] = "Items by line 項目の表"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:F1")
    ws.row_dimensions[1].height = 24
    for c, t in enumerate(("No", "項目", "数", "Price", "Amount", "日"), start=1):
        cell = ws.cell(row=2, column=c, value=t)
        cell.font, cell.fill, cell.border = Font(bold=True), head, box
        cell.alignment = Alignment(horizontal="center")
    for i in range(1, 89):
        r = i + 2
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=f"Item {i:03d} 行")
        ws.cell(row=r, column=3, value=i % 7 + 1)
        ws.cell(row=r, column=4, value=120 + i * 3).number_format = "#,##0"
        ws.cell(row=r, column=5, value=(i % 7 + 1) * (120 + i * 3)).number_format = "#,##0"
        ws.cell(row=r, column=6, value=datetime.date(2026, 9, 1) + datetime.timedelta(days=i)).number_format = "yyyy-mm-dd"
    ws["H3"] = "outside the print area"
    ws["H3"].fill = PatternFill("solid", fgColor="FFC7CE")
    ws.print_area = "A1:F90"
    ws.print_title_rows = "1:2"
    ws.page_setup.paperSize = 9
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.9, bottom=0.8, header=0.4, footer=0.35)
    ws.print_options.gridLines = True
    ws.row_breaks.append(Break(id=30))
    ws.oddHeader.left.text = "&B本書&B report"
    ws.oddHeader.center.text = "&A"
    ws.oddHeader.center.size = 14
    ws.oddHeader.right.text = "&D &F"
    ws.oddFooter.center.text = "&P / &N"
    ws.oddFooter.right.text = "Line one\nline two &&"
    ws.oddFooter.right.size = 8

    # a wide table on A4 turned: title columns, a break between columns,
    # the pages across first, scaled, with the headings
    ws = wb.create_sheet("Wide")
    ws.column_dimensions["A"].width = 12
    ws["A1"] = "行"
    ws["A1"].font = Font(bold=True)
    for c in range(2, 15):
        ws.cell(row=1, column=c, value=f"C{c - 1}").font = Font(bold=True)
        ws.cell(row=1, column=c).border = Border(bottom=medium)
    for r in range(2, 42):
        ws.cell(row=r, column=1, value=f"Row {r - 1}").fill = head
        for c in range(2, 15):
            ws.cell(row=r, column=c, value=(r - 1) * 100 + c - 1)
    ws.print_title_cols = "A:A"
    ws.print_title_rows = "1:1"
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = "landscape"
    ws.page_setup.scale = 150
    ws.page_setup.pageOrder = "overThenDown"
    ws.print_options.headings = True
    ws.print_options.horizontalCentered = True
    ws.col_breaks.append(Break(id=7))
    ws.oddFooter.center.text = "Page &P of &N"

    # more columns than a page is wide, fitted to one page across
    ws = wb.create_sheet("Fit")
    for c in range(1, 13):
        ws.column_dimensions[chr(64 + c)].width = 13
        ws.cell(row=1, column=c, value=f"Column {c}").font = Font(bold=True)
        ws.cell(row=1, column=c).fill = head
    for r in range(2, 62):
        for c in range(1, 13):
            ws.cell(row=r, column=c, value=r * c)
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = 1

    # two print areas, each a page; page numbers from 5; another header on
    # the first page; centered both ways. The borders around an area are
    # whole, and the fill beside it is not printed.
    ws = wb.create_sheet("Areas")
    for r in range(1, 7):
        for c in (1, 2, 3, 5, 6, 7):
            cell = ws.cell(row=r, column=c, value=f"{chr(64 + c)}{r}")
            cell.border = Border(left=medium, right=medium, top=medium, bottom=medium)
        ws.cell(row=r, column=4, value="no").fill = PatternFill("solid", fgColor="FFC7CE")
    ws.print_area = ["A1:C6", "E1:G6"]
    ws.page_setup.paperSize = 11
    ws.page_setup.firstPageNumber = 5
    ws.page_setup.useFirstPageNumber = True
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = True
    ws.HeaderFooter.differentFirst = True
    ws.firstHeader.center.text = "First page"
    ws.oddHeader.center.text = "Other pages"
    ws.oddFooter.left.text = "&P"

    # no print area: what the sheet uses, with the text that flows out of
    # its last column
    ws = wb.create_sheet("Used")
    ws["B2"] = "From B2"
    ws["C4"] = "This text flows out of the last column that the sheet uses, over the empty cells."
    ws.merge_cells("B6:C7")
    ws["B6"] = "merged"
    ws["B6"].alignment = Alignment(horizontal="center", vertical="center")
    ws["B6"].border = box
    ws.print_options.gridLines = True
    ws.page_setup.paperSize = 13

    path = os.path.join(OUT, "print.xlsx")
    wb.save(path)
    fix_core(path)


def forms():
    """form-fixed.xlsx, form-rows.xlsx and form-pages.xlsx: templates, whose
    cells hold tags that data replaces (xlsx.Fill): a form of one page, a
    table that grows by a row for each item, and a slip of ten rows a page
    that is copied for each page of items."""
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill("solid", fgColor="DDEBF7")

    def labels(ws, row, texts):
        for c, t in enumerate(texts, start=1):
            cell = ws.cell(row=row, column=c, value=t)
            cell.font, cell.fill, cell.border = Font(bold=True), head, box
            cell.alignment = Alignment(horizontal="center")

    # a form of one page: values of every kind
    wb = Workbook()
    props(wb, "Form of one page")
    ws = wb.active
    ws.title = "Form"
    for col, w in zip("ABCD", (8, 14, 9, 9)):
        ws.column_dimensions[col].width = w
    ws["A1"] = "明書 Certificate"
    ws["A1"].font = Font(bold=True, size=16)
    ws.merge_cells("A1:D1")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 27
    rows = [("No", "{{no}}", None), ("日", "{{date}}", "yyyy/mm/dd"), ("Name", "{{customer.name}} さま", None),
            ("Amount", "{{amount}}", "#,##0"), ("Total", '{{total | fmt:"#,##0.00"}}', None),
            ("Paid", "{{#paid}}■{{/paid}}{{^paid}}□{{/paid}} paid", None), ("Memo", "{{memo}}", None)]
    for i, (label, value, fmt) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=i, column=2, value=value)
        cell.border = Border(bottom=thin)
        if fmt:
            cell.number_format = fmt
    ws["C6"] = "=B6*1.1"
    ws["C6"].number_format = "#,##0.0"
    ws["D6"] = "=IF(B6>10000,\"large\",\"small\")"
    ws["C4"] = "｛｛date ｜ fmt：“yyyy-mm”｝｝"
    ws.merge_cells("B9:D10")
    ws["B9"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A11"] = "Photo"
    ws["A11"].font = Font(bold=True)
    ws["B11"] = "{{%photo}}"
    ws.merge_cells("B11:C13")
    for r in range(11, 14):
        for c in (2, 3):
            ws.cell(row=r, column=c).border = box
    # a QR code and a barcode of values: drawn in the cells, as pictures are
    ws["D11"] = "{{%url | qr}}"
    ws.merge_cells("D11:D13")
    for r in range(11, 14):
        ws.cell(row=r, column=4).border = box
    ws["A14"] = "Tags: {{#tags}}{{.}}{{^_last}}, {{/_last}}{{/tags}}"
    ws["A15"] = "{{missing}}"
    ws["D15"] = "\\{{not a tag}}"
    ws["A16"] = "JAN"
    ws["A16"].font = Font(bold=True)
    ws["B16"] = "{{%jan | barcode:jan}}"
    ws.merge_cells("B16:D17")
    ws.print_area = "A1:D17"
    ws.page_setup.paperSize = 11
    ws.oddFooter.center.text = "&P / &N"

    path = os.path.join(OUT, "form-fixed.xlsx")
    wb.save(path)
    fix_core(path)

    # a table that grows: a row for each item, and the formulas under it
    wb = Workbook()
    props(wb, "Table that grows")
    ws = wb.active
    ws.title = "Invoice"
    for col, w in zip("ABCDE", (6, 20, 8, 10, 12)):
        ws.column_dimensions[col].width = w
    ws["A1"] = "Invoice {{no}}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:E1")
    ws.row_dimensions[1].height = 24
    ws["A2"] = "To: {{customer}}"
    labels(ws, 3, ("No", "項目", "数", "Price", "Amount"))
    ws["A4"] = "{{#items}}{{_index}}"
    ws["B4"] = "{{name}}"
    ws["C4"] = "{{qty}}"
    ws["D4"] = "{{price}}{{/items}}"
    ws["D4"].number_format = "#,##0"
    ws["E4"] = "=C4*D4"
    ws["E4"].number_format = "#,##0"
    for c in range(1, 6):
        ws.cell(row=4, column=c).border = box
    ws["B5"] = "{{^items}}no items{{/items}}"
    ws["D6"], ws["E6"] = "Total", "=SUM(E4:E4)"
    ws["D7"], ws["E7"] = "Tax", "=ROUND(E6*0.1,0)"
    ws["D8"], ws["E8"] = "Sum", "=E6+E7"
    ws["A8"] = "=COUNT(C4:C4)"
    for r in (6, 7, 8):
        ws.cell(row=r, column=4).font = Font(bold=True)
        ws.cell(row=r, column=5).number_format = "#,##0"
        ws.cell(row=r, column=5).border = Border(bottom=thin)
    ws["A10"] = "{{note}}"
    ws.merge_cells("A10:E10")
    ws.conditional_formatting.add("E4:E4", CellIsRule(operator="greaterThanOrEqual", formula=["1000"], font=Font(bold=True, color="C00000")))
    ws.print_area = "A1:E10"
    ws.print_title_rows = "1:3"
    ws.page_setup.paperSize = 9
    ws.row_breaks.append(Break(id=8))
    ws.oddFooter.center.text = "&P / &N"
    # a second sheet that reads the first
    ws2 = wb.create_sheet("Summary")
    ws2["A1"], ws2["B1"] = "Items", "=COUNT(Invoice!C4:C4)"
    ws2["A2"], ws2["B2"] = "Sum", "=Invoice!E8"
    ws2["A3"], ws2["B3"] = "First", "=Invoice!B4"
    ws2["A4"], ws2["B4"] = "Unknown", "=TODAY()"
    ws2.print_area = "A1:B4"
    ws2.page_setup.paperSize = 11

    path = os.path.join(OUT, "form-rows.xlsx")
    wb.save(path)
    fix_core(path)

    # a slip of ten rows: the data gives the rows of each page, and what
    # is carried over
    wb = Workbook()
    props(wb, "Slip of ten rows a page")
    ws = wb.active
    ws.title = "Slip"
    for col, w in zip("ABCD", (5, 18, 7, 11)):
        ws.column_dimensions[col].width = w
    ws["A1"] = "{{#pages | pages}}Delivery {{slip}}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.row_dimensions[1].height = 24
    ws["D1"] = "{{_page}} / {{_pages}}"
    ws["D1"].alignment = Alignment(horizontal="right")
    ws["A2"] = "To: {{customer}}"
    ws["C2"], ws["D2"] = "From", "{{carry_in}}"
    ws["D2"].number_format = "#,##0"
    labels(ws, 3, ("No", "項目", "数", "Amount"))
    ws["A4"] = "{{#rows}}{{no}}"
    ws["B4"] = "{{name}}{{#_blank_first}}-- end --{{/_blank_first}}"
    ws["C4"] = "{{qty}}"
    ws["D4"] = "{{amount}}"
    ws["A13"] = "{{/rows}}"
    for r in range(4, 14):
        for c in range(1, 5):
            ws.cell(row=r, column=c).border = box
        ws.cell(row=r, column=4).number_format = "#,##0"
    ws["C14"], ws["D14"] = "Page", "=SUM(D4:D13)"
    ws["C15"], ws["D15"] = "To", "{{carry_out}}"
    ws["A15"] = "{{#_last}}Total {{total | fmt:\"#,##0\"}}{{/_last}}"
    ws["A15"].font = Font(bold=True)
    for r in (14, 15):
        ws.cell(row=r, column=4).number_format = "#,##0"
        ws.cell(row=r, column=4).border = Border(bottom=thin)
    ws["A16"] = "{{/pages}}"
    ws.print_area = "A1:D16"
    ws.page_setup.paperSize = 11
    ws.oddFooter.right.text = "&P"

    path = os.path.join(OUT, "form-pages.xlsx")
    wb.save(path)
    fix_core(path)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    if sys.argv[1:]:
        for name in sys.argv[1:]:
            globals()[name]()
    else:
        basic()
        features()
        math()
        printing()
        forms()
    print("wrote", OUT)
