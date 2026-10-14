"""Build the demo's editable Excel form templates.

Run with Python 3 and openpyxl. These are designed as business documents,
separately from the converter's feature and regression fixtures.
"""

from datetime import datetime
from io import BytesIO
from pathlib import Path
from copy import copy
import argparse
import re
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import quote_sheetname


OUT = Path(__file__).resolve().parent / "samples"
FIXED = datetime(2026, 10, 14, 0, 0, 0)
FAMILY = "M PLUS 1p"
NAVY, TEXT, MUTED, RULE, PALE = "243B53", "263342", "667685", "CBD5DF", "F3F6F9"
MONEY = '#,##0" 円"'
LINE = Side(style="thin", color=RULE)
STRONG = Side(style="medium", color=NAVY)
STYLES = ("standard", "modern")


def font(size=10, *, bold=False, color=TEXT):
    return Font(name=FAMILY, size=size, bold=bold, color=color)


def put(ws, address, text, *, size=10, bold=False, color=TEXT,
        align="left", wrap=False, fmt=None, fill=None, bottom=None):
    cell = ws[address]
    cell.value = text
    cell.font = font(size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap,
                               indent=0)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if bottom:
        cell.border = Border(bottom=bottom)
    return cell


def span(ws, cell_range, text, **kwargs):
    ws.merge_cells(cell_range)
    cell = put(ws, cell_range.split(":")[0], text, **kwargs)
    if kwargs.get("fill") or kwargs.get("bottom"):
        for row in ws[cell_range]:
            for c in row:
                if kwargs.get("fill"):
                    c.fill = PatternFill("solid", fgColor=kwargs["fill"])
                if kwargs.get("bottom"):
                    c.border = Border(bottom=kwargs["bottom"])
    return cell


def height(ws, **rows):
    for row, value in rows.items():
        ws.row_dimensions[int(row)].height = value


def new(title, last):
    wb = Workbook()
    wb.properties.title = title
    wb.properties.creator = "BDF STUDIO"
    wb.properties.lastModifiedBy = "BDF STUDIO"
    wb.properties.created = wb.properties.modified = FIXED
    ws = wb.active
    ws.title = title
    wb._fonts[0] = font(10)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.sheet_properties.outlinePr.summaryRight = False
    for col, width in zip("ABCDEFGH", (6.5, 13, 13, 13, 8.5, 11, 11, 13)):
        ws.column_dimensions[col].width = width
    for row in range(1, last + 1):
        ws.row_dimensions[row].height = 20
        for col in range(1, 9):
            cell = ws.cell(row=row, column=col)
            cell.font = font()
            cell.alignment = Alignment(vertical="center")
    ws.print_area = f"A1:H{last}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=15 / 25.4, right=15 / 25.4,
                                 top=15 / 25.4, bottom=15 / 25.4,
                                 header=7 / 25.4, footer=7 / 25.4)
    ws.print_options.horizontalCentered = True
    ws.oddFooter.left.text = "BDF STUDIO"
    ws.oddFooter.left.font = FAMILY
    ws.oddFooter.left.size = 8
    ws.oddFooter.left.color = MUTED
    ws.oddFooter.right.text = "&P / &N"
    ws.oddFooter.right.font = FAMILY
    ws.oddFooter.right.size = 8
    ws.oddFooter.right.color = MUTED
    return wb, ws


def heading(ws, title, *, no="no", customer="customer", logo=False, pages=False):
    height(ws, **{"1": 34, "2": 8, "3": 20, "4": 20, "5": 20,
                  "6": 18, "7": 14, "8": 28, "9": 18, "10": 12})
    span(ws, "A1:D1", ("{{#pages | pages}}" if pages else "") + title,
         size=24, bold=True, color=NAVY)
    if logo:
        span(ws, "G1:G1", "{{%photo}}")
        span(ws, "H1:H1", "BDF STUDIO", size=11, bold=True, color=NAVY,
             align="right")
    else:
        span(ws, "E1:H1", "BDF STUDIO", size=12, bold=True, color=NAVY,
             align="right")
    span(ws, "A3:D3", f"文書番号   {{{{{no}}}}}", size=9.5, color=MUTED)
    span(ws, "E3:H3", "〒150-0001 東京都渋谷区神宮前1-2-3", size=9, color=MUTED,
         align="right")
    span(ws, "A4:D4", 'Date   {{date | fmt:"yyyy/mm/dd"}}', size=9.5, color=MUTED)
    span(ws, "E4:H4", "TEL 03-5550-0140", size=9, color=MUTED, align="right")
    span(ws, "E5:H5", "hello@bdf.example", size=9, color=MUTED, align="right")
    span(ws, "A6:H6", "", bottom=LINE)
    span(ws, "A8:H8", f"{{{{{customer}}}}} 御中", size=14, bold=True, color=NAVY)


def table_head(ws, row, cols):
    ws.row_dimensions[row].height = 27
    for cell_range, text, alignment in cols:
        span(ws, cell_range.format(r=row), text, size=9.5, bold=True, color=NAVY,
             align=alignment, fill=PALE, bottom=LINE)


def frame(ws, address, *, strong=False):
    """Outline one logical cell, including a merged field, without inner rules."""
    cells = ws[address]
    side = Side(style="medium" if strong else "thin", color="222222")
    for r, row in enumerate(cells):
        for c, cell in enumerate(row):
            cell.border = Border(top=side if r == 0 else Side(),
                                 bottom=side if r == len(cells) - 1 else Side(),
                                 left=side if c == 0 else Side(),
                                 right=side if c == len(row) - 1 else Side())


def standardize(wb, shape):
    """A conventional Japanese business document, sharing the same fields.

    The modern layout uses a left title, tinted bands and horizontal rules.
    The standard layout separates a centred title from the sender, uses
    monochrome type, and outlines every logical cell of the detail tables.
    """
    ws = wb.active
    title = ws["A1"].value
    photo = shape == "fixed"
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row == 1:
            ws.unmerge_cells(str(merged))
    for col in range(1, 9):
        ws.cell(row=1, column=col).value = None
    span(ws, "A1:H1", title, size=24, bold=True, color="111111", align="center")
    for address in ("E3:H3", "E4:H4", "E5:H5"):
        ws.unmerge_cells(address)
    for row in range(3, 6):
        for col in range(5, 9):
            ws.cell(row=row, column=col).value = None
    if photo:
        span(ws, "G3:G3", "{{%photo}}")
        span(ws, "H3:H3", "BDF STUDIO", size=11, bold=True, align="right")
    else:
        span(ws, "E3:H3", "BDF STUDIO", size=11, bold=True, align="right")
    span(ws, "E4:H4", "〒150-0001 東京都渋谷区神宮前1-2-3",
         size=9, align="right")
    span(ws, "E5:H5", "TEL 03-5550-0140   hello@bdf.example",
         size=8.5, align="right")
    for row in ws:
        for cell in row:
            black = copy(cell.font)
            black.color = "111111" if black.bold else "222222"
            cell.font = black
            cell.fill = PatternFill(fill_type=None)
            old = cell.border
            edges = {}
            for name in ("left", "right", "top", "bottom"):
                side = getattr(old, name)
                edges[name] = Side(style=side.style, color="333333") if side and side.style else Side()
            cell.border = Border(**edges)
    for part in (ws.oddFooter.left, ws.oddFooter.right):
        part.color = "444444"
    if shape == "fixed":
        for address in ("A12:F12", "G12:H12", "A13:F13", "G13:H13", "A14:H14",
                        "E16:F16", "G16:H16", "E17:F17", "G17:H17",
                        "A22:F22", "A23:F25"):
            frame(ws, address)
        frame(ws, "E18:F18", strong=True)
        frame(ws, "G18:H18", strong=True)
    elif shape == "rows":
        for row in (15, 16):
            for part in ("A{r}:A{r}", "B{r}:E{r}", "F{r}:F{r}",
                         "G{r}:G{r}", "H{r}:H{r}"):
                frame(ws, part.format(r=row))
        for row in (18, 19, 20):
            frame(ws, f"F{row}:G{row}", strong=row == 20)
            frame(ws, f"H{row}:H{row}", strong=row == 20)
        frame(ws, "A11:D11", strong=True)
        frame(ws, "E11:H11", strong=True)
        frame(ws, "A12:H12")
    elif shape == "pages":
        for row in range(13, 24):
            for part in ("A{r}:A{r}", "B{r}:F{r}", "G{r}:G{r}", "H{r}:H{r}"):
                frame(ws, part.format(r=row))
        for row in (25, 26, 27):
            frame(ws, f"F{row}:G{row}", strong=row == 27)
            frame(ws, f"H{row}:H{row}", strong=row == 27)


def save(wb, shape, style):
    """Fix package dates as well as core properties for reproducible output."""
    if style == "standard":
        standardize(wb, shape)
    name = f"form-{shape}{'-modern' if style == 'modern' else ''}.xlsx"
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / name
    wb.save(path)
    with zipfile.ZipFile(path) as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    data = BytesIO()
    with zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as dest:
        for info, content in entries:
            if info.filename == "docProps/core.xml":
                content = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*",
                    rb"\g<1>" + FIXED.isoformat().encode() + b"Z", content)
            info.date_time = FIXED.timetuple()[:6]
            dest.writestr(info, content)
    path.write_bytes(data.getvalue())
    print(path.relative_to(OUT.parent))


def fixed(style):
    wb, ws = new("明細書", 35)
    heading(ws, "明細書", customer="customer.name", logo=True)
    span(ws, "A9:H9", "下記のとおり、ご注文の明細をご案内申し上げます。", size=9.5,
         color=MUTED)
    table_head(ws, 12, [("A{r}:F{r}", "品名・仕様", "left"),
                        ("G{r}:H{r}", "金額", "right")])
    height(ws, **{"11": 12, "13": 38, "14": 25, "15": 16,
                  "16": 23, "17": 23, "18": 34, "19": 18,
                  "20": 20, "21": 15, "22": 26, "23": 26,
                  "24": 24, "25": 24, "26": 12, "27": 23,
                  "28": 25, "29": 25, "30": 25, "31": 13,
                  "32": 18, "33": 18, "34": 18, "35": 18})
    span(ws, "A13:F13", "{{product}}", size=12, bold=True)
    span(ws, "G13:H13", "{{amount}}", size=11, align="right", fmt=MONEY)
    span(ws, "A14:H14", "{{#tags}}{{.}}{{^_last}}   /   {{/_last}}{{/tags}}",
         size=9.5, color=MUTED, bottom=LINE)
    span(ws, "E16:F16", "小計", size=9.5, color=MUTED)
    span(ws, "G16:H16", "{{amount}}", align="right", fmt=MONEY)
    span(ws, "E17:F17", "消費税（10%）", size=9.5, color=MUTED)
    span(ws, "G17:H17", "=G18-G16", align="right", fmt=MONEY)
    span(ws, "E18:F18", "総額", size=11, bold=True, color=NAVY, fill=PALE,
         bottom=STRONG)
    span(ws, "G18:H18", "{{total}}", size=17, bold=True, color=NAVY,
         align="right", fmt=MONEY, fill=PALE, bottom=STRONG)
    span(ws, "A20:H20", "お支払い状況   {{#paid}}入金済み{{/paid}}{{^paid}}未入金{{/paid}}",
         size=9.5, color=MUTED)
    span(ws, "A22:F22", "備考", size=10, bold=True, color=NAVY)
    span(ws, "A23:F25", "{{memo}}", size=10, wrap=True)
    span(ws, "G23:H25", "{{%url | qr}}")
    span(ws, "G26:H26", "注文ページ", size=8, align="center", color=MUTED)
    span(ws, "A27:H27", "商品コード", size=9.5, bold=True, color=NAVY)
    span(ws, "A28:D30", "{{%jan | barcode:jan}}")
    span(ws, "A32:H32", "", bottom=LINE)
    span(ws, "A34:H34", "ご注文ありがとうございます。ご不明な点は上記窓口までお問い合わせください。",
         size=9, color=MUTED)
    save(wb, "fixed", style)


def rows(style):
    wb, ws = new("請求書", 31)
    heading(ws, "請求書")
    span(ws, "A9:H9", "下記のとおり、ご請求申し上げます。", size=9.5, color=MUTED)
    height(ws, **{"11": 28, "12": 26, "13": 16, "14": 14,
                  "16": 28, "17": 12, "18": 20, "19": 20,
                  "20": 32, "21": 12, "22": 22, "23": 20,
                  "24": 20, "25": 20, "26": 12, "27": 22,
                  "28": 22, "29": 22, "30": 10, "31": 18})
    span(ws, "A11:D11", "ご請求金額", size=10, bold=True, color=NAVY, fill=PALE)
    span(ws, "E11:H11", "{{total}}", size=20, bold=True, color=NAVY,
         align="right", fmt=MONEY, fill=PALE)
    span(ws, "A12:H12", 'お支払期限   {{due | fmt:"yyyy/mm/dd"}}', size=10,
         color=MUTED, fill=PALE, bottom=LINE)
    table_head(ws, 15, [("A{r}:A{r}", "No.", "center"),
                        ("B{r}:E{r}", "品名", "left"),
                        ("F{r}:F{r}", "数量", "right"),
                        ("G{r}:G{r}", "単価", "right"),
                        ("H{r}:H{r}", "金額", "right")])
    put(ws, "A16", "{{#items}}{{_index}}", size=9.5, color=MUTED,
        align="center", bottom=LINE)
    span(ws, "B16:E16", "{{name}}", bottom=LINE)
    put(ws, "F16", "{{qty}}", align="right", fmt="#,##0", bottom=LINE)
    put(ws, "G16", "{{price}}", align="right", fmt="#,##0", bottom=LINE)
    put(ws, "H16", "{{amount}}{{/items}}", align="right", fmt="#,##0", bottom=LINE)
    span(ws, "F18:G18", "小計", size=9.5, color=MUTED)
    put(ws, "H18", "{{subtotal}}", align="right", fmt="#,##0")
    span(ws, "F19:G19", "消費税（10%）", size=9.5, color=MUTED)
    put(ws, "H19", "{{tax}}", align="right", fmt="#,##0")
    span(ws, "F20:G20", "総額", size=11, bold=True, color=NAVY, fill=PALE,
         bottom=STRONG)
    put(ws, "H20", "{{total}}", size=13, bold=True, color=NAVY,
        align="right", fmt="#,##0", fill=PALE, bottom=STRONG)
    span(ws, "A22:H22", "お振込先", size=10, bold=True, color=NAVY, bottom=LINE)
    span(ws, "A23:H23", "みらい銀行   渋谷支店   普通 1234567", size=9.5)
    span(ws, "A24:H24", "口座名義   ビーディーエフスタジオ", size=9.5)
    span(ws, "A25:H25", "お支払期限までに上記口座へお振り込みください。",
         size=9, color=MUTED)
    span(ws, "A27:H27", "備考", size=10, bold=True, color=NAVY, bottom=LINE)
    span(ws, "A28:H29", "{{note}}", size=9.5, wrap=True)
    span(ws, "A31:H31", "いつもお引き立てを賜り、誠にありがとうございます。", size=9,
         color=MUTED)
    ws.print_title_rows = "15:15"
    # Repeat only the table heading while detail rows continue, and move
    # complete totals/payment/note sections rather than cutting them in two.
    sheet = quote_sheetname(ws.title)
    ws.defined_names.add(DefinedName("_bdf.Print_TitleScope",
                                    attr_text=f"{sheet}!$A$15:$H$16"))
    ws.defined_names.add(DefinedName("_bdf.Print_KeepTogether", attr_text=",".join(
        f"{sheet}!{ref}" for ref in ("$F$18:$H$20", "$A$22:$H$25", "$A$27:$H$31"))))
    ws.freeze_panes = "B16"
    save(wb, "rows", style)


def pages(style):
    wb, ws = new("納品書", 34)
    heading(ws, "納品書", no="slip", pages=True)
    span(ws, "A9:F9", "下記のとおり、納品いたしました。", size=9.5, color=MUTED)
    span(ws, "G9:H9", "{{_page}} / {{_pages}}", size=9.5, color=MUTED, align="right")
    height(ws, **{"11": 26, "12": 16, "14": 31, "15": 31,
                  "16": 31, "17": 31, "18": 31, "19": 31,
                  "20": 31, "21": 31, "22": 31, "23": 31,
                  "24": 14, "25": 24, "26": 24, "27": 34,
                  "28": 10, "29": 10, "30": 12, "31": 12,
                  "32": 18, "33": 18, "34": 1})
    span(ws, "A11:F11", "前頁繰越", size=9.5, color=MUTED)
    span(ws, "G11:H11", "{{carry_in}}", size=10, align="right", fmt=MONEY)
    table_head(ws, 13, [("A{r}:A{r}", "No.", "center"),
                        ("B{r}:F{r}", "品名", "left"),
                        ("G{r}:G{r}", "数量", "right"),
                        ("H{r}:H{r}", "金額", "right")])
    put(ws, "A14", "{{#rows}}{{no}}", size=9.5, color=MUTED,
        align="center", bottom=LINE)
    span(ws, "B14:F14", "{{name}}", bottom=LINE)
    put(ws, "G14", "{{qty}}", align="right", fmt="#,##0", bottom=LINE)
    put(ws, "H14", "{{amount}}", align="right", fmt="#,##0", bottom=LINE)
    for row in range(15, 24):
        span(ws, f"B{row}:F{row}", "", bottom=LINE)
        for col in "AGH":
            put(ws, f"{col}{row}", "", align="center" if col == "A" else "right",
                fmt="#,##0", bottom=LINE)
    # Ten rows are slots. Closing the section in the last slot keeps the
    # row height and its rules when a page contains fewer than ten items.
    put(ws, "A23", "{{/rows}}", size=9.5, color=MUTED,
        align="center", bottom=LINE)
    span(ws, "F25:G25", "当頁計", size=9.5, color=MUTED)
    put(ws, "H25", "{{page_total}}", align="right", fmt="#,##0")
    span(ws, "F26:G26", "{{#_last}}累計{{/_last}}{{^_last}}次頁繰越{{/_last}}",
         size=9.5, color=MUTED)
    put(ws, "H26", "{{carry_out}}", align="right", fmt="#,##0", bottom=LINE)
    span(ws, "F27:G27", "{{#_last}}総額{{/_last}}{{^_last}}次頁へ続きます{{/_last}}",
         size=11, bold=True, color=NAVY,
         fill=PALE, bottom=STRONG)
    put(ws, "H27", "{{#_last}}{{total | fmt:\"#,##0\"}}{{/_last}}", size=13,
        bold=True, color=NAVY, align="right", fill=PALE, bottom=STRONG)
    span(ws, "A30:H30", "", bottom=LINE)
    span(ws, "A32:H32", "数量・内容をご確認のうえ、相違がございましたらご連絡ください。",
         size=9, color=MUTED)
    span(ws, "A33:H33", "このたびはご注文いただき、誠にありがとうございます。", size=9,
         color=MUTED)
    put(ws, "A34", "{{/pages}}", size=1)
    save(wb, "pages", style)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--theme", choices=STYLES, help="write one style; default: both")
    args = parser.parse_args()
    for style in (args.theme,) if args.theme else STYLES:
        fixed(style)
        rows(style)
        pages(style)
