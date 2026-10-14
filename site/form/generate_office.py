"""Build the public form demo's Word and PowerPoint templates.

Run with Python 3 and python-pptx. These are deliberately separate from the
converter fixtures: the website demonstrates readable business documents,
whereas test/docx and test/pptx exercise every supported template feature.
"""
from __future__ import annotations

import importlib.util
import io
from pathlib import Path
import re
import sys
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent / "samples"
FONT = "M PLUS 1p"
NAVY, TEXT, MUTED, RULE, PALE = "243B53", "263342", "667685", "CBD5DF", "F3F6F9"
THEME = "modern"
W, H, M = 595, 842, 42
CW = W - 2 * M
ISSUER = ("BDF STUDIO", "〒150-0001 東京都渋谷区神宮前1-2-3", "TEL 03-5550-0140", "hello@bdf.example")
BANK = "みらい銀行 渋谷支店　普通 1234567　ビーディーエフスタジオ"


def load_helper(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


# Importing the fixture's pure OOXML helpers must not dirty its directory.
sys.dont_write_bytecode = True
D = load_helper("demo_docx_helpers", ROOT / "test/docx/gen.py")


def deterministic_zip(path):
    """Normalize both package timestamps and Office core timestamps."""
    with zipfile.ZipFile(path) as archive:
        parts = [(info.filename, archive.read(info.filename)) for info in archive.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts:
            if name == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(created|modified)[^>]*>)[^<]*", rb"\g<1>2026-10-14T00:00:00Z", data)
            info = zipfile.ZipInfo(name, (2026, 10, 14, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, data)


def logo_bytes():
    path = OUT / "brand-mark.png"
    return path.read_bytes() if path.exists() else D.png(16, 16, lambda x, y: (36, 59, 83))


def dr(text, size=10, bold=False, color=None):
    color = TEXT if color is None else color
    return D.run(text, f'<w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:eastAsia="{FONT}"/>'
                 f'<w:sz w:val="{round(size * 2)}"/><w:color w:val="{color}"/>' + ('<w:b/>' if bold else ''))


def dp(text="", size=10, bold=False, color=None, align="left", before=0, after=0, keep=False, raw=False, style=None):
    ppr = (f'<w:spacing w:before="{round(before * 20)}" w:after="{round(after * 20)}" w:line="{round(size * 1.5 * 20)}" w:lineRule="atLeast"/>'
           f'<w:jc w:val="{align}"/>' + ('<w:keepNext/>' if keep else ''))
    return D.para(text if raw else dr(text, size, bold, color), ppr, style)


def spacer(height=10):
    return D.para(dr("", 1), f'<w:spacing w:before="0" w:after="0" w:line="{round(height * 20)}" w:lineRule="exact"/>')


def dc(text, width, align="left", fill=None, bold=False, size=10, raw=False, span=1):
    shade = f'<w:shd w:val="clear" w:fill="{fill}"/>' if fill else ""
    spans = f'<w:gridSpan w:val="{span}"/>' if span > 1 else ""
    body = text if raw else dp(text, size, bold, align=align)
    return f'<w:tc><w:tcPr><w:tcW w:w="{round(width * 20)}" w:type="dxa"/><w:vAlign w:val="center"/>{shade}{spans}</w:tcPr>{body}</w:tc>'


def dt(widths, rows, header=0, heights=None, borders=True, align="left"):
    width = sum(widths)
    grid = "".join(f'<w:gridCol w:w="{round(x * 20)}"/>' for x in widths)
    border = "".join(f'<w:{edge} w:val="single" w:sz="4" w:color="{RULE}"/>' if borders and (THEME == "standard" or edge in ("top", "bottom", "insideH")) else f'<w:{edge} w:val="nil"/>' for edge in ("top", "left", "bottom", "right", "insideH", "insideV"))
    out = []
    for i, cells in enumerate(rows):
        pr = '<w:tblHeader/>' if i < header else ""
        if heights:
            pr += f'<w:trHeight w:val="{round(heights[i] * 20)}" w:hRule="atLeast"/>'
        out.append(f'<w:tr><w:trPr>{pr}<w:cantSplit/></w:trPr>{"".join(cells)}</w:tr>')
    return (f'<w:tbl><w:tblPr><w:tblW w:w="{round(width * 20)}" w:type="dxa"/><w:tblLayout w:type="fixed"/><w:jc w:val="{align}"/>'
            f'<w:tblBorders>{border}</w:tblBorders><w:tblCellMar><w:top w:w="90" w:type="dxa"/>'
            '<w:left w:w="120" w:type="dxa"/><w:bottom w:w="90" w:type="dxa"/><w:right w:w="120" w:type="dxa"/></w:tblCellMar>'
            f'</w:tblPr><w:tblGrid>{grid}</w:tblGrid>{"".join(out)}</w:tbl>')


def dhead(title, number, customer, intro, photo=False):
    brand = D.inline_pic("rIdPhoto", 24 * 12700, 24 * 12700, 11, "BDF brand mark", "{{%photo}}") + dr("  BDF STUDIO", 11, True, NAVY) if photo else dr("BDF STUDIO", 11, True, NAVY)
    meta = dp(f"No. {{{{{number}}}}}", 9, color=MUTED, align="right") + dp('Date {{date | fmt:"yyyy-mm-dd"}}', 9, color=MUTED, align="right")
    titleblock = dp(title, 24, True, NAVY, after=10, style="Title") + D.para(brand, '<w:spacing w:after="0"/>')
    out = [dt([CW * .58, CW * .42], [[dc(titleblock, CW * .58, raw=True), dc(meta, CW * .42, raw=True)]], borders=False), spacer(20)]
    if THEME == "standard":
        out = [dp(title, 24, True, align="center", after=14, style="Title"),
               dt([CW * .58, CW * .42], [[dc(D.para(brand), CW * .58, raw=True), dc(meta, CW * .42, raw=True)]], borders=False), spacer(8)]
    issuer = "".join(dp(line, 9 if i else 11, i == 0, MUTED if i else NAVY, align="right") for i, line in enumerate(ISSUER))
    recipient = dp("お届け先" if title == "納品書" else "宛先", 9, color=MUTED, after=5) + dp(f"{{{{{customer}}}}} 御中", 14, True, after=10) + dp(intro, 9.5)
    out.extend([dt([CW * .55, CW * .45], [[dc(recipient, CW * .55, raw=True), dc(issuer, CW * .45, raw=True)]], borders=False), spacer(22)])
    return out


def dsummary(labels, values, total_last=True):
    label_parts, value_parts = [], []
    for i, (label, value) in enumerate(zip(labels, values)):
        total = total_last and i == len(labels) - 1
        height = 35 if total else 28
        spacing = f'<w:spacing w:before="0" w:after="0" w:line="{height * 20}" w:lineRule="exact"/>'
        label_parts.append(D.para(dr(label, 10, total), spacing))
        value_parts.append(D.para(dr(value, 16 if total else 10, total), spacing + '<w:jc w:val="right"/>'))
    # A single unbreakable row keeps the subtotal, tax and total together.
    if THEME == "standard":
        row = [dc("".join(label_parts), 120, raw=True), dc("".join(value_parts), 110, raw=True)]
        return dt([120, 110], [row], align="right")
    row = [dc("", CW-230), dc("".join(label_parts), 120, fill=PALE if total_last else None, raw=True), dc("".join(value_parts), 110, fill=PALE if total_last else None, raw=True)]
    return dt([CW-230, 120, 110], [row], borders=False)


def dfooter():
    tabs = f'<w:tabs><w:tab w:val="right" w:pos="{round(CW * 20)}"/></w:tabs>'
    line = f'<w:pBdr><w:top w:val="single" w:sz="4" w:space="10" w:color="{RULE}"/></w:pBdr>'
    return D.hdr_ftr("ftr", D.para(dr("BDF STUDIO", 8, color=MUTED) + dr("\t", 8) + D.field("PAGE", "1", f'<w:sz w:val="16"/><w:color w:val="{MUTED}"/>') + dr(" / ", 8, color=MUTED) + D.field("NUMPAGES", "1", '<w:sz w:val="16"/>'), tabs + line))


def write_docx(shape, body, title, photo=False):
    refs = '<w:footerReference w:type="default" r:id="rIdFooter"/>'
    section = (f'<w:sectPr>{refs}<w:pgSz w:w="{W * 20}" w:h="{H * 20}"/><w:pgMar w:top="{M * 20}" w:right="{M * 20}" '
               f'w:bottom="{M * 20}" w:left="{M * 20}" w:header="400" w:footer="460"/></w:sectPr>')
    document = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {D.NS}><w:body>{"".join(body)}{section}</w:body></w:document>'
    styles = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:styles {D.NS}><w:docDefaults>'
              f'<w:rPrDefault><w:rPr><w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:eastAsia="{FONT}"/><w:sz w:val="20"/><w:color w:val="{TEXT}"/></w:rPr></w:rPrDefault>'
              '<w:pPrDefault><w:pPr><w:spacing w:before="0" w:after="0"/></w:pPr></w:pPrDefault></w:docDefaults>'
              '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>'
              f'<w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:rPr><w:sz w:val="48"/><w:b/><w:color w:val="{NAVY}"/></w:rPr></w:style></w:styles>')
    parts = {"word/styles.xml": styles, "word/numbering.xml": D.NUMBERING, "word/settings.xml": D.settings(),
             "word/theme/theme1.xml": D.THEME, "word/footer1.xml": dfooter(), "docProps/core.xml": D.core(title, "BDF STUDIO")}
    rels = {"rIdStyles": ("styles", "styles.xml"), "rIdFooter": ("footer", "footer1.xml"), "rIdSettings": ("settings", "settings.xml"), "rIdTheme": ("theme", "theme/theme1.xml")}
    if photo:
        parts["word/media/brand.png"] = logo_bytes()
        rels["rIdPhoto"] = ("image", "media/brand.png")
    types = '<Override PartName="/word/footer1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>'
    suffix = "-modern" if THEME == "modern" else ""
    path = OUT / f"form-{shape}{suffix}.docx"
    D.package(path, document, parts, rels, types)
    deterministic_zip(path)


def docx_fixed():
    body = dhead("明細書", "no", "customer.name", "下記のとおり、商品のお取引明細をご案内いたします。", True)
    body += [dp("お取引明細", 11, True, NAVY, after=8)]
    widths = [CW - 130, 130]
    body += [dt(widths, [[dc("商品・仕様", widths[0], fill=PALE, bold=True), dc("金額（税抜）", widths[1], "right", PALE, True)],
                        [dc(dp("{{product}}", 11, True, after=5) + dp("{{#tags}}{{.}}{{^_last}} / {{/_last}}{{/tags}}", 9, color=MUTED), widths[0], raw=True), dc('{{amount | fmt:"¥#,##0"}}', widths[1], "right")]], header=1, heights=[28, 64]), spacer(14)]
    body += [dsummary(["総額（税込）"], ['{{total | fmt:"¥#,##0"}}']), spacer(18), dp("お支払い状況", 9, color=MUTED, after=4),
             dp("{{#paid}}お支払い済み{{/paid}}{{^paid}}お支払い待ち{{/paid}}", 10, True), spacer(20),
             dp("備考", 10, True, NAVY, after=7), dp("{{memo}}", 10), spacer(32)]
    qr = D.inline_pic("rIdPhoto", 44 * 12700, 44 * 12700, 12, "Product details QR", "{{%url | qr}}")
    jan = D.inline_pic("rIdPhoto", 140 * 12700, 44 * 12700, 13, "JAN barcode", "{{%jan | barcode:jan}}")
    body += [dt([CW - 180, 180], [[dc(dp("商品情報", 9, color=MUTED, after=6) + dp(qr, raw=True), CW - 180, raw=True), dc(dp("JANコード", 9, color=MUTED, align="right", after=6) + dp(jan, raw=True, align="right"), 180, raw=True)]], borders=False)]
    write_docx("fixed", body, "明細書", True)


def docx_rows():
    body = dhead("請求書", "no", "customer", "下記のとおりご請求申し上げます。")
    body += [dsummary(["ご請求総額（税込）"], ['{{total | fmt:"¥#,##0"}}']), spacer(12), dp('お支払期限　{{due | fmt:"yyyy-mm-dd"}}', 10, after=16)]
    widths = [32, CW - 32 - 50 - 82 - 94, 50, 82, 94]
    header = [dc(t, w, "right" if i > 1 else "left", PALE, True, 9.5) for i, (t, w) in enumerate(zip(("No.", "品名", "数量", "単価", "金額"), widths))]
    row = [dc("{{#items}}{{_index}}", widths[0], "center", size=9.5), dc("{{name}}", widths[1], size=9.5), dc("{{qty}}", widths[2], "right", size=9.5), dc('{{price | fmt:"#,##0"}}', widths[3], "right", size=9.5), dc('{{amount | fmt:"#,##0"}}{{/items}}', widths[4], "right", size=9.5)]
    body += [dt(widths, [header, row], header=1, heights=[28, 31]), spacer(12), dsummary(["小計", "消費税（10%）", "総額（税込）"], ['{{subtotal | fmt:"¥#,##0"}}', '{{tax | fmt:"¥#,##0"}}', '{{total | fmt:"¥#,##0"}}']), spacer(22),
             dp("お振込先", 10, True, NAVY, after=6), dp(BANK, 9.5), spacer(16), dp("備考", 10, True, NAVY, after=6), dp("{{note}}", 9.5)]
    write_docx("rows", body, "請求書")


def docx_pages():
    body = [dp("{{#pages | pages}}")]
    body += dhead("納品書", "slip", "customer", "下記の商品を納品いたしました。")
    body += [dp('前頁繰越　{{carry_in | fmt:"¥#,##0"}}', 10, align="right", after=12)]
    widths = [34, CW - 34 - 60 - 112, 60, 112]
    rows = [[dc(t, w, "right" if i > 1 else "left", PALE, True, 9.5) for i, (t, w) in enumerate(zip(("No.", "品名", "数量", "金額"), widths))],
            [dc("{{#rows}}{{no}}", widths[0], "center", size=9.5), dc("{{name}}{{#_blank_first}}以下余白{{/_blank_first}}", widths[1], size=9.5), dc("{{qty}}", widths[2], "right", size=9.5), dc('{{amount | fmt:"#,##0"}}', widths[3], "right", size=9.5)]]
    for i in range(9):
        rows.append([dc("", widths[0]), dc("", widths[1]), dc("", widths[2]), dc("{{/rows}}" if i == 8 else "", widths[3])])
    body += [dt(widths, rows, header=1, heights=[28] + [31] * 10), spacer(12), dsummary(["当頁計", "累計金額"], ['{{page_total | fmt:"¥#,##0"}}', '{{carry_out | fmt:"¥#,##0"}}'], False)]
    if THEME == "standard":
        body += [dp("{{#_last}}"), dsummary(["総額（税込）"], ['{{total | fmt:"¥#,##0"}}']), dp("{{/_last}}"), spacer(12)]
    else:
        body += [dp('{{#_last}}総額　{{total | fmt:"¥#,##0"}}{{/_last}}', 16, True, NAVY, align="right", before=8, after=18)]
    body += [dp("備考", 10, True, NAVY, after=5), dp("お受け取り後、数量・商品の状態をご確認ください。", 9.5), dp("{{/pages}}")]
    write_docx("pages", body, "納品書")


def pt_color(value):
    return RGBColor.from_string(value)


def ppt_text(slide, text, x, y, width, height, size=10, bold=False, color=None, align=PP_ALIGN.LEFT):
    color = TEXT if color is None else color
    shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(width), Pt(height))
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.35
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = pt_color(color)
        r._r.get_or_add_rPr().append(r._r.makeelement(qn("a:ea"), {"typeface": FONT}))
    return shape


def ppt_line(slide, y):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(M), Pt(y), Pt(W-M), Pt(y))
    shape.line.width = Pt(.6)
    shape.line.color.rgb = pt_color(RULE)
    shape._element.spPr.append(shape._element.makeelement(qn("a:effectLst"), {}))


def ppt_pic(slide, descr, x, y, width, height):
    shape = slide.shapes.add_picture(io.BytesIO(logo_bytes()), Pt(x), Pt(y), Pt(width), Pt(height))
    shape._element.nvPicPr.cNvPr.set("descr", descr)
    return shape


def ppt_box(slide, x, y, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(width), Pt(height))
    shape.fill.background()
    shape.line.color.rgb = pt_color(RULE)
    shape.line.width = Pt(.8)
    shape._element.spPr.append(shape._element.makeelement(qn("a:effectLst"), {}))
    return shape


def ppt_head(slide, title, number, customer, intro, photo=False, page_repeat=False):
    prefix = "{{#pages | pages}}" if page_repeat else ""
    centered = THEME == "standard"
    ppt_text(slide, prefix + title, M, 44, CW if centered else 250, 34, 24, True, NAVY, PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT)
    ppt_text(slide, f"No. {{{{{number}}}}}\nDate {{{{date | fmt:\"yyyy-mm-dd\"}}}}", W-M-190, 88 if centered else 48, 190, 34, 9, color=MUTED, align=PP_ALIGN.RIGHT)
    if photo:
        ppt_pic(slide, "{{%photo}}", M, 94, 24, 24)
    ppt_text(slide, "BDF STUDIO", M+32 if photo else M, 96, 210, 20, 11, True, NAVY)
    if not centered:
        ppt_line(slide, 130)
    ppt_text(slide, "お届け先" if title == "納品書" else "宛先", M, 151, 250, 16, 9, color=MUTED)
    ppt_text(slide, f"{{{{{customer}}}}} 御中", M, 172, 298, 24, 14, True)
    ppt_text(slide, intro, M, 208, 305, 26, 9.5)
    for i, line in enumerate(ISSUER):
        ppt_text(slide, line, 325, 151 + i * 17, CW - 283, 17, 11 if i == 0 else 8.5, i == 0, NAVY if i == 0 else MUTED, PP_ALIGN.RIGHT)


def ppt_footer(slide, pages=False):
    ppt_line(slide, 790)
    ppt_text(slide, "BDF STUDIO", M, 800, 260, 14, 8, color=MUTED)
    tag = "{{_page}} / {{_pages}}{{/pages}}" if pages else "{{_slide}} / {{_slides}}"
    ppt_text(slide, tag, W-M-100, 800, 100, 14, 8, color=MUTED, align=PP_ALIGN.RIGHT)


def ppt_table(slide, y, widths, rows, row_height=31, headers=1, totals=0):
    shape = slide.shapes.add_table(len(rows), len(widths), Pt(M), Pt(y), Pt(sum(widths)), Pt(row_height * len(rows)))
    t = shape.table
    t.first_row = t.horz_banding = t.vert_banding = False
    for c, width in enumerate(widths):
        t.columns[c].width = Pt(width)
    for i, row in enumerate(rows):
        t.rows[i].height = Pt(row_height)
        for j, text in enumerate(row):
            cell = t.cell(i, j)
            cell.margin_left = cell.margin_right = Pt(7)
            cell.margin_top = cell.margin_bottom = Pt(5)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = pt_color(PALE if i < headers or (totals and i >= len(rows)-totals) else "FFFFFF")
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if j >= (2 if len(widths) > 2 else 1) else PP_ALIGN.LEFT
            if j == 0 and len(widths) > 2:
                p.alignment = PP_ALIGN.CENTER
            p.space_before = p.space_after = Pt(0)
            r = p.add_run()
            r.text = text
            r.font.name = FONT
            r.font.size = Pt(9.5)
            r.font.bold = i < headers or (totals and i == len(rows)-1)
            r.font.color.rgb = pt_color(TEXT)
            r._r.get_or_add_rPr().append(r._r.makeelement(qn("a:ea"), {"typeface": FONT}))
            pr = cell._tc.get_or_add_tcPr()
            for edge in ("L", "R", "T", "B"):
                ln = pr.makeelement(qn(f"a:ln{edge}"), {"w": "6350"})
                if THEME == "modern" and edge in ("L", "R"):
                    ln.append(ln.makeelement(qn("a:noFill"), {}))
                    pr.append(ln)
                    continue
                fill = ln.makeelement(qn("a:solidFill"), {})
                fill.append(fill.makeelement(qn("a:srgbClr"), {"val": RULE}))
                ln.append(fill)
                ln.append(ln.makeelement(qn("a:prstDash"), {"val": "solid"}))
                pr.append(ln)
    return t


def ppt_summary(slide, label, value, y):
    if THEME == "standard":
        ppt_box(slide, M, y-3, CW, 44)
    ppt_text(slide, label, M+9 if THEME == "standard" else M, y+7, 230, 19, 10, True, NAVY)
    ppt_text(slide, value, 270, y, CW-237 if THEME == "standard" else CW-228, 34, 23, True, NAVY, PP_ALIGN.RIGHT)
    if THEME != "standard":
        ppt_line(slide, y+40)


def new_ppt():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(W), Pt(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return prs, slide


def save_ppt(prs, shape, title):
    prs.core_properties.title = title
    prs.core_properties.author = "BDF STUDIO"
    suffix = "-modern" if THEME == "modern" else ""
    path = OUT / f"form-{shape}{suffix}.pptx"
    prs.save(path)
    deterministic_zip(path)


def pptx_fixed():
    prs, slide = new_ppt()
    ppt_head(slide, "明細書", "no", "customer.name", "下記のとおり、商品のお取引明細をご案内いたします。", True)
    ppt_text(slide, "お取引明細", M, 262, 220, 18, 11, True, NAVY)
    ppt_table(slide, 290, [CW-140, 140], [["商品・仕様", "金額（税抜）"], ["{{product}}\n{{#tags}}{{.}}{{^_last}} / {{/_last}}{{/tags}}", '{{amount | fmt:"¥#,##0"}}']], 40)
    ppt_summary(slide, "総額（税込）", '{{total | fmt:"¥#,##0"}}', 405)
    ppt_text(slide, "お支払い状況", M, 470, 220, 18, 9, color=MUTED)
    ppt_text(slide, "{{#paid}}お支払い済み{{/paid}}{{^paid}}お支払い待ち{{/paid}}", M, 491, 260, 20, 10, True)
    ppt_text(slide, "備考", M, 538, 220, 18, 10, True, NAVY)
    ppt_text(slide, "{{memo}}", M, 564, CW, 62, 10)
    ppt_text(slide, "商品情報", M, 667, 130, 16, 9, color=MUTED)
    ppt_pic(slide, "{{%url | qr}}", M, 691, 44, 44)
    ppt_text(slide, "JANコード", W-M-160, 667, 160, 16, 9, color=MUTED, align=PP_ALIGN.RIGHT)
    ppt_pic(slide, "{{%jan | barcode:jan}}", W-M-140, 691, 140, 44)
    ppt_footer(slide)
    save_ppt(prs, "fixed", "明細書")


def pptx_rows():
    prs, slide = new_ppt()
    ppt_head(slide, "請求書", "no", "customer", "下記のとおりご請求申し上げます。")
    ppt_summary(slide, "ご請求総額（税込）", '{{total | fmt:"¥#,##0"}}', 254)
    ppt_text(slide, 'お支払期限　{{due | fmt:"yyyy-mm-dd"}}', M, 310, CW, 19, 10)
    widths = [32, CW-32-50-82-94, 50, 82, 94]
    ppt_table(slide, 351, widths, [["No.", "品名", "数量", "単価", "金額"],
                                 ["{{#items}}{{_index}}", "{{name}}", "{{qty}}", '{{price | fmt:"#,##0"}}', '{{amount | fmt:"#,##0"}}{{/items}}']], 28)
    if THEME == "standard":
        ppt_box(slide, 314, 587, 239, 76)
    for i, (label, value) in enumerate((("小計", '{{subtotal | fmt:"¥#,##0"}}'), ("消費税（10%）", '{{tax | fmt:"¥#,##0"}}'), ("総額（税込）", '{{total | fmt:"¥#,##0"}}'))):
        ppt_text(slide, label, 323, 594+i*22, 120, 20, 10, i == 2)
        ppt_text(slide, value, 443, 594+i*22, 103 if THEME == "standard" else 110, 20, 12 if i == 2 else 10, i == 2, align=PP_ALIGN.RIGHT)
    ppt_text(slide, "お振込先", M, 668, 220, 18, 10, True, NAVY)
    ppt_text(slide, BANK, M, 690, CW, 20, 9.5)
    ppt_text(slide, "備考", M, 726, 220, 17, 10, True, NAVY)
    ppt_text(slide, "{{note}}", M, 748, CW, 30, 9.5)
    ppt_footer(slide)
    save_ppt(prs, "rows", "請求書")


def pptx_pages():
    prs, slide = new_ppt()
    ppt_head(slide, "納品書", "slip", "customer", "下記の商品を納品いたしました。", page_repeat=True)
    ppt_text(slide, '前頁繰越　{{carry_in | fmt:"¥#,##0"}}', M, 253, CW, 19, 10, align=PP_ALIGN.RIGHT)
    rows = [["No.", "品名", "数量", "金額"], ["{{#rows}}{{no}}", "{{name}}{{#_blank_first}}以下余白{{/_blank_first}}", "{{qty}}", '{{amount | fmt:"#,##0"}}']]
    rows += [["", "", "", "{{/rows}}" if i == 8 else ""] for i in range(9)]
    ppt_table(slide, 286, [34, CW-34-60-112, 60, 112], rows, 31)
    ppt_text(slide, '当頁計　{{page_total | fmt:"¥#,##0"}}', M, 640, CW, 19, 10, align=PP_ALIGN.RIGHT)
    ppt_text(slide, '累計金額　{{carry_out | fmt:"¥#,##0"}}', M, 664, CW, 19, 10, align=PP_ALIGN.RIGHT)
    if THEME == "standard":
        ppt_box(slide, 314, 689, 239, 38)
        ppt_text(slide, '総額　{{total | fmt:"¥#,##0"}}', 323, 696, 223, 28, 19, True, NAVY, PP_ALIGN.RIGHT)
    else:
        ppt_text(slide, '{{#_last}}総額　{{total | fmt:"¥#,##0"}}{{/_last}}', M, 696, CW, 28, 19, True, NAVY, PP_ALIGN.RIGHT)
    ppt_text(slide, "備考", M, 739, 220, 16, 10, True, NAVY)
    ppt_text(slide, "お受け取り後、数量・商品の状態をご確認ください。", M, 761, CW, 18, 9.5)
    ppt_footer(slide, True)
    save_ppt(prs, "pages", "納品書")


def main():
    global THEME, NAVY, TEXT, MUTED, RULE, PALE
    OUT.mkdir(exist_ok=True)
    for THEME in ("standard", "modern"):
        NAVY, TEXT, MUTED, RULE, PALE = ("000000", "000000", "444444", "222222", "FFFFFF") if THEME == "standard" else ("243B53", "263342", "667685", "CBD5DF", "F3F6F9")
        for builder in (docx_fixed, docx_rows, docx_pages, pptx_fixed, pptx_rows, pptx_pages):
            builder()


if __name__ == "__main__":
    main()
